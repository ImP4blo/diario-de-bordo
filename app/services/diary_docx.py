"""Documento Word do diário de bordo final (o que vai para a empresa).

Formato:
    Diário de Bordo
    Nome: Fulano · Período: 26/10/2026 a 30/10/2026
    Segunda-feira, 26/10/2026
    | Início | Fim | Atividade | Descrição |
"""

from datetime import date

from docx.shared import Cm, Pt

from app.schemas import DiaryDay
from app.services.diary import NO_RECORD
from app.services.docx_utils import CellStyle, add_table, new_document, to_bytes
from app.services.transcript import format_day

# Soma 17cm: a largura útil de uma folha A4 com margens de 2cm
COLUMN_WIDTHS = (Cm(1.6), Cm(1.6), Cm(4.6), Cm(9.2))
HEADERS = ("Início", "Fim", "Atividade", "Descrição")


def format_period(start: date, end: date) -> str:
    return f"{start:%d/%m/%Y}" if start == end else f"{start:%d/%m/%Y} a {end:%d/%m/%Y}"


def diary_filename(start: date, end: date) -> str:
    """'diario-de-bordo_2026-10-26_a_2026-10-30.docx'"""
    period = f"{start}" if start == end else f"{start}_a_{end}"
    return f"diario-de-bordo_{period}.docx"


def render_diary_docx(person: str, start: date, end: date, days: list[DiaryDay]) -> bytes:
    doc = new_document()
    doc.add_heading("Diário de Bordo", level=0)

    info = doc.add_paragraph()
    info.add_run("Nome: ").bold = True
    info.add_run(person)
    info.add_run("     Período: ").bold = True
    info.add_run(format_period(start, end))
    info.paragraph_format.space_after = Pt(12)

    for diary in days:
        doc.add_heading(format_day(diary.day), level=2)
        rows = [
            (
                (f"{b.inicio:%H:%M}", f"{b.fim:%H:%M}", b.atividade, b.descricao),
                # "Sem registro" discreto, para não parecer uma atividade
                CellStyle(italic=True, muted=True) if b.atividade == NO_RECORD else CellStyle(),
            )
            for b in diary.blocks
        ]
        add_table(doc, HEADERS, COLUMN_WIDTHS, rows)

    return to_bytes(doc)
