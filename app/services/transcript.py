"""Exportação da transcrição bruta: texto (.txt / .md) ou Word (.docx).

Serve para revisar o dia ou colar num chat de IA e montar o diário manualmente.

Organização: primeiro `collect()` junta os dados (dias e linhas); depois cada
formato só "desenha" esses dados. Assim texto e Word nunca divergem.
"""

from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from itertools import groupby
from typing import Literal
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.models import Entry, EntryKind
from app.schemas import RoutineData
from app.services.days import as_utc

ExportFormat = Literal["txt", "docx", "md"]

# Nomes fixos em vez do "locale" do sistema, que no Windows nem sempre está em pt-BR
WEEKDAYS = [
    "Segunda-feira",
    "Terça-feira",
    "Quarta-feira",
    "Quinta-feira",
    "Sexta-feira",
    "Sábado",
    "Domingo",
]

KIND_LABELS = {
    EntryKind.SPEECH: "Fala",
    EntryKind.NOTE: "ANOTAÇÃO",
    EntryKind.PAUSE: "Pausa",
    EntryKind.CONFIRM: "Confirmação",
    EntryKind.ACTIVITY_CHANGE: "Nova atividade",
    EntryKind.ACTIVITY_END: "Fim da atividade",
}
ROUTINE_LABEL = "Rotina"

PROMPT_INTRO = "Como usar: cole este arquivo inteiro num chat de IA junto com o pedido abaixo."
PROMPT = (
    "Monte meu diário de bordo a partir da transcrição abaixo. Devolva uma tabela por dia "
    "com as colunas Início | Fim | Atividade | Descrição. Regras: português profissional, "
    "frases curtas; use só os horários reais da transcrição, sem inventar atividades; "
    "agrupe falas da mesma atividade num bloco; ignore conversa irrelevante e ruído; "
    "linhas marcadas como ANOTAÇÃO têm prioridade; inclua os blocos [Rotina] como estão; "
    'marque lacunas dentro do expediente como "sem registro".'
)

TITLE = "Transcrição: Diário de Bordo"


@dataclass
class Line:
    time: str  # "10:05" ou "12:02–13:05"
    label: str  # "Fala", "ANOTAÇÃO", "Rotina"...
    text: str

    def __str__(self) -> str:
        return f"{self.time} [{self.label}] {self.text}".rstrip()


@dataclass
class Transcript:
    person: str
    period: str
    workday: str | None
    generated_at: str
    days: list[tuple[date, list[Line]]] = field(default_factory=list)


def format_day(day: date) -> str:
    """'Segunda-feira, 26/10/2026'"""
    return f"{WEEKDAYS[day.weekday()]}, {day:%d/%m/%Y}"


def _entry_line(entry: Entry, tz: ZoneInfo) -> Line:
    time = f"{as_utc(entry.ts).astimezone(tz):%H:%M}"
    if entry.ends_at:
        time += f"–{as_utc(entry.ends_at).astimezone(tz):%H:%M}"
    label = KIND_LABELS.get(entry.kind, entry.kind)
    return Line(time, label, entry.text or (entry.activity or ""))


def _day_lines(
    day: date, entries: Iterable[Entry], routine: RoutineData | None, tz: ZoneInfo
) -> list[Line]:
    """Linhas do dia: entradas + blocos da rotina (só em dias úteis), por horário."""
    timed = [(as_utc(e.ts).astimezone(tz).time(), _entry_line(e, tz)) for e in entries]
    if routine and day.weekday() < 5:
        timed += [
            (b.start, Line(f"{b.start:%H:%M}–{b.end:%H:%M}", ROUTINE_LABEL, b.activity))
            for b in routine.blocks
        ]
    # sort é "estável": no mesmo horário, a ordem original se mantém
    timed.sort(key=lambda item: item[0])
    return [line for _, line in timed]


def collect(
    entries: Iterable[Entry],
    start: date,
    end: date,
    person: str,
    tz: ZoneInfo,
    routine: RoutineData | None = None,
) -> Transcript:
    """Junta tudo o que vai no arquivo. `entries` deve vir ordenado por horário."""
    period = f"{start:%d/%m/%Y}" if start == end else f"{start:%d/%m/%Y} a {end:%d/%m/%Y}"
    workday = f"{routine.work_start:%H:%M} às {routine.work_end:%H:%M}" if routine else None
    transcript = Transcript(person, period, workday, f"{datetime.now(tz):%d/%m/%Y %H:%M}")
    for day, day_entries in groupby(entries, key=lambda e: e.local_date):
        transcript.days.append((day, _day_lines(day, day_entries, routine, tz)))
    return transcript


def _info_lines(t: Transcript) -> list[str]:
    lines = [f"Pessoa: {t.person}", f"Período: {t.period}"]
    if t.workday:
        lines.append(f"Expediente: {t.workday}")
    lines.append(f"Gerado em: {t.generated_at}")
    return lines


# --- Texto (.txt / .md) ------------------------------------------------------


def render_text(t: Transcript, fmt: Literal["txt", "md"] = "txt") -> str:
    md = fmt == "md"
    lines = [f"# {TITLE}" if md else TITLE.upper(), "", *_info_lines(t), ""]
    if md:
        lines += [f"> **{PROMPT_INTRO}**", ">", f"> {PROMPT}"]
    else:
        lines += [PROMPT_INTRO, "", PROMPT]

    for day, day_lines in t.days:
        lines += ["", f"## {format_day(day)}" if md else f"== {format_day(day)} ==", ""]
        lines += [f"- {line}" if md else str(line) for line in day_lines]

    if not t.days:
        lines += ["", "Nenhum registro no período."]
    return "\n".join(lines) + "\n"


# --- Word (.docx) --------------------------------------------------------------

HEADER_BG = "2F6FDE"
MUTED = RGBColor(0x64, 0x6B, 0x76)
COLUMN_WIDTHS = (Cm(2.6), Cm(3.2), Cm(11.2))  # Horário | Tipo | Texto (A4 com margem 2cm)


def _shade(cell, hex_color: str) -> None:
    """python-docx não tem atalho para cor de fundo da célula: editamos o XML direto."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def render_docx(t: Transcript) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading(TITLE, level=0)
    for info in _info_lines(t):
        doc.add_paragraph(info).paragraph_format.space_after = Pt(0)

    prompt = doc.add_paragraph()
    prompt.paragraph_format.space_before = Pt(10)
    prompt.add_run(f"{PROMPT_INTRO} ").bold = True
    prompt.add_run(PROMPT).italic = True

    if not t.days:
        doc.add_paragraph("Nenhum registro no período.")

    for day, day_lines in t.days:
        doc.add_heading(format_day(day), level=2)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for cell, title in zip(table.rows[0].cells, ("Horário", "Tipo", "Texto"), strict=True):
            _shade(cell, HEADER_BG)
            run = cell.paragraphs[0].add_run(title)
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        for line in day_lines:
            cells = table.add_row().cells
            for cell, value in zip(cells, (line.time, line.label, line.text), strict=True):
                run = cell.paragraphs[0].add_run(value)
                if line.label == KIND_LABELS[EntryKind.NOTE]:
                    run.bold = True  # anotações em destaque
                elif line.label == ROUTINE_LABEL:
                    run.italic = True
                    run.font.color.rgb = MUTED

        # No Word, a largura precisa ser definida célula por célula
        for row in table.rows:
            for cell, width in zip(row.cells, COLUMN_WIDTHS, strict=True):
                cell.width = width

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
